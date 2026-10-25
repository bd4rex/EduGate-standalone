"""Local document parsing with source locations retained for classroom citations."""
from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from fastapi import HTTPException

PARSER_VERSION = 2
SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".json", ".html", ".htm", ".xml", ".log", ".pdf", ".docx", ".pptx"}
MAX_TEXT_CHARS = 4_000_000


@dataclass
class Section:
    text: str
    chapter: str = ""
    page: int | None = None


@dataclass
class ParsedDocument:
    sections: list[Section]
    warnings: list[str] = field(default_factory=list)


def clean_text(text: str) -> str:
    lines = [re.sub(r"[^\S\n]+", " ", line).strip() for line in text.replace("\r", "").splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _text_sections(text: str, *, page: int | None = None, chapter: str = "") -> list[Section]:
    result: list[Section] = []
    body: list[str] = []
    headings: list[str] = []
    current = chapter
    fenced = False

    def flush() -> None:
        value = clean_text("\n".join(body))
        if value:
            result.append(Section(value, current, page))
        body.clear()

    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("```") or stripped.startswith("~~~"):
            fenced = not fenced
        match = None if fenced else re.match(r"^(#{1,6})\s+(.+)$", stripped)
        chinese = not fenced and len(stripped) < 100 and re.match(r"^(?:第[一二三四五六七八九十百\d]+[章节单元]|[一二三四五六七八九十]+、)", stripped)
        if match or chinese:
            flush()
            if match:
                level, title = len(match[1]), match[2].strip()
                headings = headings[:level - 1] + [title]
                current = " / ".join(headings)
            else:
                current = stripped
            continue
        if not stripped and not fenced:
            flush()
        else:
            # Textbook separators have no retrieval value.
            if not re.fullmatch(r"[=\-_]{3,}", stripped):
                body.append(line)
    flush()
    return result


def _validate_office_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(item.file_size for item in entries) > 200 * 1024 * 1024:
                raise HTTPException(status_code=413, detail="Office 文档展开后超过解析限制")
    except zipfile.BadZipFile as error:
        raise HTTPException(status_code=400, detail="不是有效的 Word / PowerPoint 文档") from error


def parse_document(path: Path, *, max_pdf_pages: int = 500) -> ParsedDocument:
    suffix = path.suffix.lower()
    result = ParsedDocument([])
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        if len(reader.pages) > max_pdf_pages:
            raise HTTPException(status_code=413, detail=f"PDF exceeds the {max_pdf_pages}-page limit")
        blank: list[int] = []
        total = 0
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            total += len(text)
            if total > MAX_TEXT_CHARS:
                raise HTTPException(status_code=413, detail="文档文字超过解析限制，请按章节拆分")
            if not text.strip():
                blank.append(number)
            result.sections.extend(_text_sections(text, page=number))
        if blank:
            result.warnings.append(f"{len(blank)} 页未抽取到文字（页码：{', '.join(map(str, blank[:20]))}）；可能需要 OCR，图像内容尚未进入检索。")
    elif suffix == ".docx":
        _validate_office_archive(path)
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(str(path))
        heading = ""
        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, document)
                value = paragraph.text.strip()
                if paragraph.style and paragraph.style.name.lower().startswith(("heading", "标题")):
                    heading = value
                elif value:
                    result.sections.append(Section(value, heading))
            elif child.tag.endswith("}tbl"):
                table = Table(child, document)
                value = "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
                result.sections.append(Section(value, heading))
        if document.inline_shapes:
            result.warnings.append("Word 中的图片已保留在原文件中；图片、公式的视觉内容未自动识别。")
    elif suffix == ".pptx":
        _validate_office_archive(path)
        from pptx import Presentation
        presentation = Presentation(str(path))
        image_count = 0
        for number, slide in enumerate(presentation.slides, 1):
            heading = slide.shapes.title.text if slide.shapes.title else f"幻灯片 {number}"
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text)
                if shape.has_table:
                    parts.append("\n".join(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows))
                if shape.shape_type == 13:
                    image_count += 1
            value = clean_text("\n\n".join(parts))
            if value:
                result.sections.append(Section(value, heading, number))
        if image_count:
            result.warnings.append(f"{image_count} 张图片保留在原文件中，视觉内容未自动识别。")
    else:
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise HTTPException(status_code=400, detail="无法识别文本编码，请保存为 UTF-8")
        if suffix in {".html", ".htm", ".xml"}:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(text, "html.parser")
            for element in soup(["script", "style", "noscript"]):
                element.decompose()
            for element in soup.find_all(re.compile(r"^h[1-6]$")):
                element.replace_with(f"\n{'#' * int(element.name[1])} {element.get_text(' ', strip=True)}\n")
            text = soup.get_text("\n")
        result.sections = _text_sections(text)
    if sum(len(section.text) for section in result.sections) > MAX_TEXT_CHARS:
        raise HTTPException(status_code=413, detail="文档文字超过解析限制，请按章节拆分")
    return result


def split_sections(document: ParsedDocument, *, size: int = 350, overlap: int = 45) -> list[Section]:
    """Keep each chunk inside one page/chapter; leave room for BGE's 512-token context."""
    if size <= overlap or overlap < 0:
        raise ValueError("Invalid chunk size or overlap")
    merged: list[Section] = []
    for section in document.sections:
        text = clean_text(section.text)
        if not text:
            continue
        if merged and (merged[-1].page, merged[-1].chapter) == (section.page, section.chapter) and len(merged[-1].text) + len(text) + 2 <= size:
            merged[-1].text += "\n\n" + text
        else:
            merged.append(Section(text, section.chapter, section.page))
    chunks: list[Section] = []
    for section in merged:
        start = 0
        while start < len(section.text):
            end = min(start + size, len(section.text))
            if end < len(section.text):
                boundaries = [match.end() for match in re.finditer(r"[。！？；\n]", section.text[start:end])]
                if boundaries and boundaries[-1] >= size // 2:
                    end = start + boundaries[-1]
            chunks.append(Section(section.text[start:end].strip(), section.chapter, section.page))
            if end == len(section.text):
                break
            start = max(start + 1, end - overlap)
    return chunks
