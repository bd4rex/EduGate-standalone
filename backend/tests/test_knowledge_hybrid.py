from __future__ import annotations

import asyncio
import json
import os
import shutil
import sqlite3
from pathlib import Path

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app.knowledge import KnowledgeStore, KnowledgeMetadata, KnowledgeScope, tokenize
from app.knowledge_documents import parse_document, split_sections


def add_text(store: KnowledgeStore, filename: str, text: str, source: str = "general"):
    (store.source_directory(source) / filename).write_text(text, encoding="utf-8")
    result = store.scan_source(source)
    assert not result.errors
    return next(file for file in store.list_files(source) if file.filename == filename)


@pytest.fixture
def store(tmp_path):
    value = KnowledgeStore(str(tmp_path / "knowledge.sqlite3"), str(tmp_path / "files"), semantic_enabled=False)
    yield value
    value.close()


def test_chinese_title_identifier_and_long_query(store):
    file = add_text(store, "量化.md", "# 图像编码\n把连续变化的取值映射到有限的离散等级。\n\nNAT 将内网地址转换为公网地址。")
    assert store.search("general", "量化")[0].file_id == file.id
    query = "老师请帮我详细解释这些相关的内容" * 20 + " NAT"
    assert "nat" in tokenize(query)
    assert store.search("general", query)
    assert store.search("general", "量化")[0].chapter == "图像编码"


def test_scope_versions_chapters_and_teacher_materials_are_enforced(store):
    first = add_text(store, "2025.md", "# 子网\n子网掩码区分网络部分和主机部分。")
    second = add_text(store, "2026.md", "# IPv6\nIPv6 地址有128位。")
    answers = add_text(store, "答案.md", "子网掩码题目的标准答案是选项C。")
    store.update_file_metadata(first.id, KnowledgeMetadata(subject="信息技术", edition="2025"))
    store.update_file_metadata(second.id, KnowledgeMetadata(subject="信息技术", edition="2026"))
    store.update_file_metadata(answers.id, KnowledgeMetadata(audience="teacher"))
    hits = store.search("general", "子网掩码")
    assert hits and all(hit.file_id != answers.id for hit in hits)
    assert not store.search("general", "IPv6", scope=KnowledgeScope(edition="2025"))
    assert not store.search("general", "子网掩码", scope=KnowledgeScope(chapter="IPv6"))
    assert not store.search("general", "子网掩码", scope=KnowledgeScope(file_ids=[second.id]))
    with pytest.raises(HTTPException):
        store.validate_scope("general", KnowledgeScope(file_ids=[answers.id]))
    store.update_file_metadata(first.id, KnowledgeMetadata(edition="x' OR true --"))
    assert store.search("general", "子网掩码", scope=KnowledgeScope(edition="x' OR true --"))[0].file_id == first.id


def test_update_delete_and_restore_rebuild_derived_index(store, tmp_path):
    file = add_text(store, "lesson.txt", "alpha original evidence")
    old = store.search("general", "alpha")[0]
    snapshot = tmp_path / "backup.sqlite3"
    with sqlite3.connect(snapshot) as target, store._connect() as source:
        source.backup(target)
    add_text(store, "lesson.txt", "bravo replacement evidence")
    assert not store.search("general", "alpha")
    assert store.search("general", "bravo")
    with pytest.raises(HTTPException):
        store.citation(old.chunk_id, "general", KnowledgeScope(), content_hash=old.content_hash)
    store.delete_file(file.id)
    assert not store.search("general", "bravo")
    # Restore authoritative SQLite into another app; a pre-existing cache must not win.
    restored = KnowledgeStore(str(snapshot), str(store.storage_dir), index_dir=str(store.index.directory), semantic_enabled=False)
    try:
        assert restored.search("general", "alpha")
        assert not restored.search("general", "bravo")
    finally:
        restored.close()


def test_parse_failure_does_not_serve_old_version(store):
    add_text(store, "lesson.txt", "alpha original evidence")
    assert store.search("general", "alpha")
    (store.source_directory("general") / "lesson.txt").write_bytes(b"")
    scan = store.scan_source("general")
    assert scan.errors and store.list_files()[0].index_error
    assert not store.search("general", "alpha")
    add_text(store, "lesson.txt", "bravo repaired evidence")
    assert store.search("general", "bravo")


@pytest.mark.parametrize("change", ["delete", "teacher", "replace"])
def test_concurrent_change_cannot_return_stale_snapshot_text(store, monkeypatch, change):
    file = add_text(store, "lesson.txt", "alpha original evidence")
    store.warm_index()
    original = store.index.search
    def race(*args, **kwargs):
        results = original(*args, **kwargs)
        if change == "delete":
            store.delete_file(file.id)
        elif change == "teacher":
            store.update_file_metadata(file.id, KnowledgeMetadata(audience="teacher"))
        else:
            add_text(store, "lesson.txt", "bravo replacement evidence")
        return results
    monkeypatch.setattr(store.index, "search", race)
    assert store.search("general", "alpha") == []


def test_word_and_slides_keep_headings_tables_and_locations(tmp_path):
    from docx import Document
    from pptx import Presentation
    word = Document()
    word.add_heading("网络地址", level=1)
    word.add_paragraph("子网掩码用于区分网络部分与主机部分。")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "IPv6"
    table.cell(0, 1).text = "128位"
    path = tmp_path / "lesson.docx"
    word.save(path)
    parsed = split_sections(parse_document(path))
    assert any("IPv6 | 128位" in part.text for part in parsed)
    assert all(part.chapter == "网络地址" and part.page is None for part in parsed)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "NAT"
    slide.placeholders[1].text = "多台设备共用一个公网地址。"
    path = tmp_path / "lesson.pptx"
    deck.save(path)
    section = parse_document(path).sections[0]
    assert section.page == 1 and section.chapter == "NAT"


def test_pdf_page_locations_and_unreadable_page_warning(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject("/Type"):NameObject("/Font"),NameObject("/Subtype"):NameObject("/Type1"),NameObject("/BaseFont"):NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):DictionaryObject({NameObject("/F1"):writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 14 Tf 40 700 Td (IPv6 addresses contain 128 bits.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_blank_page(width=600, height=800)
    path = tmp_path / "text-and-scan.pdf"
    writer.write(path)
    parsed = parse_document(path)
    assert parsed.sections[0].page == 1 and "128" in parsed.sections[0].text
    assert parsed.warnings and "2" in parsed.warnings[0]


def test_followup_query_keeps_previous_question_without_client_scope():
    from app.chat_service import _retrieval_query
    from app.schemas import ChatRequest
    request = ChatRequest(messages=[{"role":"user","content":"NAT有什么作用？"},{"role":"assistant","content":"说明"},{"role":"user","content":"你能再举个例子吗？"}])
    assert "NAT" in _retrieval_query(request) and "例子" in _retrieval_query(request)
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        ChatRequest(messages=[{"role":"user","content":"test"}], knowledge_scope={"file_ids":[]})


@pytest.fixture
def api_client(store, monkeypatch):
    from app import state
    from app.main import app
    from app.routers import knowledge, config
    monkeypatch.setattr(state, "knowledge_store", store)
    monkeypatch.setattr(knowledge, "knowledge_store", store)
    monkeypatch.setattr(config, "knowledge_store", store)
    monkeypatch.setattr("app.dependencies._client_ip", lambda _: "127.0.0.1")
    monkeypatch.setattr(state.settings, "classroom_rate_limit", 1000)
    snapshot = state.runtime_config.data.model_copy(deep=True)
    with TestClient(app) as client:
        if not state.business_db.is_admin_initialized(state.settings.admin_username):
            state.business_db.setup_admin(username=state.settings.admin_username,password="test-admin-password",display_name="Teacher")
        state.classroom_access.start()
        state.rate_limiter._events.clear()
        teacher = {"X-Admin-Token":state.sessions.issue(state.settings.admin_username)}
        student = client.post("/classroom/join",headers={"X-Class-Token":state.classroom_access.token()}).json()
        yield client, teacher, {"X-Student-Token":student["student_token"]}
    state.runtime_config.data = snapshot
    state.runtime_config.save()


def test_api_upload_scope_preview_and_student_citation(api_client):
    from app import state
    client, teacher, student = api_client
    assert client.get("/knowledge/status",headers=student).status_code == 401
    response = client.post("/knowledge/files",headers=teacher,data={"source_id":"general"},files={"file":("lesson.md","# IPv6\nIPv6地址长度是128位。".encode(),"text/plain")})
    assert response.status_code == 200
    file = response.json()
    assert client.get(f"/knowledge/files/{file['id']}/preview",headers=teacher).json()["chunks"][0]["chapter"] == "IPv6"
    assert "attachment" in client.get(f"/knowledge/files/{file['id']}/original",headers=teacher).headers["content-disposition"]
    saved = client.put("/config/scenarios/default",headers=teacher,json={"knowledge_source_id":"general","knowledge_scope":{"file_ids":[file["id"]]}})
    assert saved.status_code == 200
    hit = client.post("/knowledge/search",headers=teacher,json={"source_id":"general","query":"IPv6"}).json()["hits"][0]
    uri = f"/knowledge/citations/{hit['chunk_id']}?version={hit['content_hash']}"
    assert client.get(uri,headers=student).status_code == 200
    client.patch(f"/knowledge/files/{file['id']}/metadata",headers=teacher,json={"audience":"teacher"})
    assert client.get(uri,headers=student).status_code == 404
    state.classroom_access.end()
    assert client.get(uri,headers=student).status_code == 403


def test_chat_strict_uses_evidence_once_and_streams_sources(api_client, store, monkeypatch):
    from app import chat_service, state
    from app.routers import chat
    client, teacher, student = api_client
    add_text(store, "IPv6.md", "IPv6地址长度是128位。")
    client.put("/config/scenarios/default",headers=teacher,json={"knowledge_source_id":"general","knowledge_strict":True})
    calls = []
    original = store.search_details
    def retrieve(*args, **kwargs):
        calls.append(args[1]); return original(*args,**kwargs)
    monkeypatch.setattr(store,"search_details",retrieve)
    async def completion(payload):
        if "check whether" in payload["messages"][0]["content"]:
            evidence = json.loads(payload["messages"][-1]["content"])
            return {"choices":[{"message":{"content":json.dumps({"supported":True,"evidence":[{"id":"K1","quote":"IPv6地址长度是128位。"}]})}}]}
        return {"choices":[{"message":{"content":"IPv6地址有128位。[K1]"}}]}
    monkeypatch.setattr(chat_service,"_chat_completion",completion)
    monkeypatch.setattr(chat,"_chat_completion",completion)
    response = client.post("/chat",headers=student,json={"messages":[{"role":"user","content":"IPv6有多少位？"}]})
    assert response.status_code == 200 and len(calls) == 1
    assert response.json()["edugate_knowledge"]["sources"][0]["filename"] == "IPv6.md"
    log = state.business_db.list_logs(1)[0]
    assert log["retrieval"]["decision"] == "supported"
    assert log["retrieval"]["chunks"][0]["file_id"]
    async def stream(payload):
        yield 'data: {"choices":[{"delta":{"content":"128位。[K1]"}}]}\n\ndata: [DONE]\n\n'.encode()
    monkeypatch.setattr(chat,"_stream_with_errors",stream)
    response = client.post("/chat/stream",headers=student,json={"messages":[{"role":"user","content":"IPv6有多少位？"}]})
    assert response.status_code == 200 and "event: knowledge" in response.text and "[DONE]" in response.text
    assert len(calls) == 2
    assert client.post("/chat",headers=student,json={"scenario_id":"other","messages":[{"role":"user","content":"question"}]}).status_code == 403


def test_strict_mode_blocks_when_no_source_is_selected(monkeypatch):
    from app import chat_service
    from app.schemas import TeachingScenario, ChatRequest
    request = ChatRequest(messages=[{"role":"user","content":"请给出答案"}])
    assert asyncio.run(chat_service._strict_miss_decision(request,TeachingScenario(knowledge_strict=True))) == (True,False)


@pytest.mark.skipif(not os.getenv("EDUGATE_TEST_KNOWLEDGE_MODEL_DIR"), reason="requires the prepared real local model; no download in unit tests")
def test_real_local_semantic_retrieval_and_model_provenance(tmp_path):
    model_dir = os.environ["EDUGATE_TEST_KNOWLEDGE_MODEL_DIR"]
    store = KnowledgeStore(str(tmp_path/"real.sqlite3"),str(tmp_path/"files"),model_dir=model_dir,semantic_enabled=True)
    try:
        assert store.embeddings.load()
        nat = add_text(store,"网络地址转换.md","# NAT\nNAT把内网的私有IP转换为公网IP，让多台设备共享一个外部地址。")
        add_text(store,"图像编码.md","# 量化\n量化把连续变化的取值转换为有限的离散等级。")
        result = store.search_details("general","我家几台电脑怎样共用一个外网号码？")
        assert result.mode == "hybrid" and result.hits[0].file_id == nat.id
        assert result.hits[0].semantic_score is not None
        assert store.status()["embedded_count"] == store.status()["chunk_count"]
    finally:
        store.close()
